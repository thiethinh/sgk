import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_highlight_box(doc, text_list, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.3)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="20" w:space="0" w:color="2B6CB0"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(3)
    if title:
        run_title = p.add_run(f"📌 {title}\n")
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(11.5)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
    for i, t in enumerate(text_list):
        if i > 0 or title:
            p2 = cell.add_paragraph()
        else:
            p2 = p
        p2.paragraph_format.line_spacing = 1.2
        p2.paragraph_format.space_after = Pt(2)
        run = p2.add_run(t)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def build_summary_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    COLOR_PRIMARY = RGBColor(0x1A, 0x36, 0x5D)     # Navy
    COLOR_SECONDARY = RGBColor(0x2B, 0x6C, 0xB0)   # Blue
    COLOR_DARK = RGBColor(0x2D, 0x37, 0x48)        # Charcoal

    # Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h = p_header.add_run("TÀI LIỆU CỐT LÕI PHỤC VỤ THUYẾT TRÌNH & THIẾT KẾ BÀI GIẢNG")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(10.5)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    p_header.paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("BẢN TÓM TẮT Ý TƯỞNG & DỰ ÁN: BÉ TẬP ĐỌC LỚP 1")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("(Dàn ý súc tích, đắt giá phục vụ biên soạn slide trình chiếu)")
    r_s.font.name = "Times New Roman"
    r_s.font.size = Pt(11)
    r_s.font.italic = True
    r_s.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    p_sub.paragraph_format.space_after = Pt(14)

    def add_section_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_bullet_item(bold_txt, regular_txt):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        rb = p.add_run(bold_txt + " ")
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(11.5)
        rb.font.bold = True
        rb.font.color.rgb = COLOR_SECONDARY
        
        rr = p.add_run(regular_txt)
        rr.font.name = "Times New Roman"
        rr.font.size = Pt(11.5)
        rr.font.color.rgb = COLOR_DARK
        return p

    # 1. TIÊU ĐỀ & THÔNG ĐIỆP
    add_section_title("1. TÊN ĐỀ TÀI & THÔNG ĐIỆP CỐT LÕI")
    create_highlight_box(
        doc,
        [
            "Tên đề tài: Bé Tập Đọc Lớp 1 – Sách Giáo Khoa Điện Tử Tương Tác Tích Hợp Từ Điển Song Ngữ.",
            "Thông điệp truyền cảm hứng: \"Chạm vào đâu, hiểu sâu và đọc chuẩn ở đó\" – Biến cuốn sách giáo khoa quen thuộc thành một người bạn đồng hành sống động, khơi nguồn niềm say mê học đọc cho con trẻ."
        ],
        title="Thông Điệp Chủ Đạo"
    )

    # 2. VẤN ĐỀ TỪ THỰC TẾ
    add_section_title("2. BỐI CẢNH & NHỮNG TRĂN TRỞ THỰC TẾ (VẤN ĐỀ)")
    add_bullet_item("Khó khăn của học sinh lớp 1:", "Bước ngoặt từ mầm non sang tiểu học; sách giấy in tĩnh thiếu âm thanh tương tác; trẻ dễ nản lòng khi gặp vần khó và sợ đọc sai trước đám đông.")
    add_bullet_item("Nỗi vất vả của thầy cô:", "Tiết dạy tập đọc giáo viên phải đọc mẫu lặp đi lặp lại nhiều lần rất mất sức; khó duy trì sự tập trung hào hứng của 30 - 40 học sinh cùng lúc.")
    add_bullet_item("Trăn trở của phụ huynh:", "Bận rộn công việc; nhiều phụ huynh lo ngại mình phát âm chưa chuẩn hoặc không rành tiếng Anh để kèm con học bài đều đặn mỗi tối.")
    add_bullet_item("Học ngoại ngữ bị tách rời:", "Cách học từ vựng tiếng Anh riêng rẽ khiến trẻ học vẹt, khó liên hệ giữa từ ngữ mẹ đẻ và từ tiếng Anh trong cuộc sống.")

    # 3. Ý TƯỞNG CỐT LÕI
    add_section_title("3. Ý TƯỞNG GIẢI PHÁP: CUỐN SÁCH \"BIẾT NÓI & BIẾT GIẢI NGHĨA\"")
    add_bullet_item("Giữ nguyên 100% sách chuẩn của Bộ GD&ĐT:", "Không thay đổi nét vẽ hay làm méo mó tranh minh họa, bảo toàn trọn vẹn sự thân thuộc của bộ sách 'Kết nối tri thức với cuộc sống'.")
    add_bullet_item("Bố cục hai phần đồng hành:", "Trang sách gốc nằm trọn bên trái, thẻ giải nghĩa từ vựng xuất hiện song song ngay bên phải mà không che mất trang sách.")
    add_bullet_item("Tương tác đa giác quan:", "Học qua trải nghiệm trực quan: Nhìn tranh minh họa – Chạm vào mặt chữ – Nghe giọng đọc mẫu – Hiểu nghĩa từ ngữ.")

    # 4. HIỆN THỰC HÓA
    add_section_title("4. SẢN PHẨM THỰC TẾ ĐÃ HOÀN THIỆN (ĐÃ LÀM ĐƯỢC GÌ?)")
    add_bullet_item("57 trang sách điện tử hoàn chỉnh:", "Số hóa toàn bộ các bài học trong học kỳ 1 của Sách giáo khoa Tiếng Việt 1; căn chỉnh vừa vặn tầm mắt nhìn trên mọi màn hình.")
    add_bullet_item("Thao tác tự nhiên như mở sách thật:", "Chạm nhẹ tay (hoặc nhấp chuột) vào chữ là nghe đọc ngay; lăn chuột nhẹ nhàng để chuyển trang êm ái, không nhấp nháy.")
    add_bullet_item("Kho giọng đọc mẫu mực (hơn 200 bản thu âm):", "Giọng tiếng Việt truyền cảm, tròn vành rõ chữ theo chuẩn mẫu giáo khoa; giọng tiếng Anh chuẩn bản ngữ.")
    add_bullet_item("Lời giải nghĩa ngây thơ, ấm áp:", "Biên soạn riêng cho lứa tuổi 6 tuổi, gắn liền với tình cảm gia đình, trường lớp (ví dụ: giải nghĩa từ 'bà', 'gia đình', 'búp bê').")
    add_bullet_item("Sổ tay từ vựng yêu thích của bé:", "Tính năng 'thả tim' lưu lại các từ hay để bé mở ra ôn luyện riêng biệt, tạo cảm giác sưu tầm huy hiệu học tập.")
    add_bullet_item("Thiết kế tối giản, an toàn cho trẻ:", "Lược bỏ mọi nút bấm phức tạp hay quảng cáo gây xao nhãng; có mảnh ghi chú nhỏ xinh chỉ dẫn bé thao tác dễ dàng.")

    # 5. LỢI ÍCH TO LỚN
    add_section_title("5. GIÁ TRỊ VÀ LỢI ÍCH TO LỚN MANG LẠI")
    add_bullet_item("Với Học sinh lớp 1:", "Xóa tan nỗi sợ đọc sai; tự tin đọc to trước lớp; hình thành tính tự giác khám phá; làm quen tiếng Anh tự nhiên theo ngữ cảnh bài học tiếng Việt.")
    add_bullet_item("Với Thầy cô giáo:", "Giáo cụ sinh động trên bảng tương tác/tivi thông minh; giảm tải việc đọc mẫu liên tục; giờ học sôi nổi; bám sát 100% tiến trình giáo án hiện hành.")
    add_bullet_item("Với Phụ huynh:", "An tâm để con tự học bài tại nhà vì có giọng đọc chuẩn mực đồng hành; tiết kiệm thời gian kèm cặp, giải tỏa áp lực học tập.")

    # 6. ĐỊNH HƯỚNG MỞ RỘNG
    add_section_title("6. HƯỚNG PHÁT TRIỂN TIẾP THEO")
    add_bullet_item("Nhận diện giọng đọc của bé:", "Tích hợp micro thông minh để bé tự đọc to và máy chấm điểm động viên, khích lệ.")
    add_bullet_item("Trò chơi ôn tập nhẹ nhàng:", "Đố vui ghép chữ với tranh minh họa, chọn từ đúng để nhận hoa điểm 10.")
    add_bullet_item("Mở rộng học liệu số toàn diện:", "Hoàn thiện tiếp Tập 2 và nhân rộng sang các khối lớp 2, 3, 4, 5.")

    # 7. THÔNG ĐIỆP KẾT LUẬN
    add_section_title("7. KẾT LUẬN")
    create_highlight_box(
        doc,
        [
            "\"Bé Tập Đọc Lớp 1\" là minh chứng cho việc công nghệ thực sự phục vụ con người – khi công nghệ được chắt lọc để trở nên vô hình, giản dị, nhường chỗ cho niềm vui học tập của con trẻ và sự tiện lợi của thầy cô trên bục giảng."
        ],
        title="Thông Điệp Kết Thúc Bài Thuyết Trình"
    )

    out_file = r"c:\Users\HINH\Downloads\sgk1\TOM_TAT_THUYET_TRINH.docx"
    doc.save(out_file)
    print(f"Generated: {out_file}")

if __name__ == "__main__":
    build_summary_doc()
